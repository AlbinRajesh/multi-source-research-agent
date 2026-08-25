"""
Word document text extraction parser.
Extracts clean text from .docx files using python-docx.
"""

import logging
from pathlib import Path

from docx import Document
from docx.opc.exceptions import PackageNotFoundError

logger = logging.getLogger(__name__)


class DocxParseError(Exception):
    """Raised when a Word document cannot be parsed."""
    pass


def parse_docx(file_path: str) -> str:
    """
    Extract all text content from a .docx file, including paragraphs and tables.

    Args:
        file_path: Path to the .docx file.

    Returns:
        Extracted text, with paragraphs and table rows separated by newlines.

    Raises:
        DocxParseError: If the file doesn't exist, isn't a valid .docx,
                         or contains no extractable text.
    """
    path = Path(file_path)

    if not path.exists():
        raise DocxParseError(f"File not found: {file_path}")

    if path.suffix.lower() != ".docx":
        raise DocxParseError(f"Not a .docx file: {file_path} "
                              f"(note: legacy .doc files are not supported)")

    try:
        doc = Document(file_path)
    except PackageNotFoundError:
        raise DocxParseError(f"Could not open .docx (corrupted or invalid): {file_path}")
    except Exception as e:
        raise DocxParseError(f"Unexpected error opening .docx: {e}")

    text_parts = []

    # Paragraphs (regular body text)
    for para in doc.paragraphs:
        stripped = para.text.strip()
        if stripped:
            text_parts.append(stripped)

    # Tables (docx stores these separately from paragraphs). Each table is
    # kept as its own block, separated by a blank line from every other
    # table/paragraph — matching pdf_parser.py's approach. Without this,
    # multiple tables' rows get concatenated into one continuous stream,
    # so downstream parsing treats table 2's first row as a continuation
    # of table 1's header, silently rejecting every row as "misaligned".
    table_blocks = []
    for table in doc.tables:
        row_lines = []
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                row_lines.append(row_text)
        if row_lines:
            table_blocks.append("\n".join(row_lines))

    prose_block = "\n".join(text_parts)
    all_blocks = [b for b in ([prose_block] + table_blocks) if b.strip()]
    full_text = "\n\n".join(all_blocks)
    if not full_text.strip():
        raise DocxParseError(
            f"No extractable text found in {file_path}. "
            f"The document may be empty or contain only images."
        )

    logger.info(f"Extracted {len(full_text)} characters from {path.name} "
                f"({len(doc.paragraphs)} paragraphs, {len(doc.tables)} table(s))")
    return full_text